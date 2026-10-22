"""Read source documents and experiment evidence without executing embedded content."""
from __future__ import annotations

import base64
import csv
import warnings as python_warnings
from io import BytesIO, StringIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader

MAX_BYTES = 20 * 1024 * 1024
MAX_CHARS = 180_000
MAX_IMAGE_PIXELS = 20_000_000
REFERENCE_SUFFIXES = {".txt", ".md", ".docx", ".pdf"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp"}


def _paragraph_text(paragraph: Paragraph) -> str:
    """Keep Office equation token text; equation layout still needs a warning."""
    if "oMath" not in paragraph._p.xml:
        return paragraph.text
    tokens = []
    for element in paragraph._p.iter():
        local = element.tag.rsplit("}", 1)[-1]
        if local == "t" and element.text:
            tokens.append(element.text)
        elif local == "tab":
            tokens.append("\t")
        elif local in {"br", "cr"}:
            tokens.append("\n")
    return "".join(tokens)


def _checked_archive(raw: bytes) -> None:
    try:
        with ZipFile(BytesIO(raw)) as archive:
            entries = archive.infolist()
            if len(entries) > 15_000 or sum(entry.file_size for entry in entries) > 150 * 1024 * 1024:
                raise ValueError("文档解压后过大，无法安全解析，请拆分后上传。")
    except BadZipFile as exc:
        raise ValueError("文件不是有效的 Office 文档。") from exc


def _image(raw: bytes, filename: str) -> dict:
    try:
        from PIL import Image
    except ImportError as exc:
        raise ValueError("图片解析需要 Pillow，请安装项目依赖后重试。") from exc
    try:
        with python_warnings.catch_warnings():
            python_warnings.simplefilter("error", Image.DecompressionBombWarning)
            with Image.open(BytesIO(raw)) as picture:
                fmt = picture.format
                if fmt not in {"PNG", "JPEG", "WEBP"}:
                    raise ValueError("图片只支持 PNG、JPEG 或 WEBP。")
                if picture.width * picture.height > MAX_IMAGE_PIXELS:
                    raise ValueError("图片不能超过 2000 万像素，请缩小后上传。")
                if getattr(picture, "n_frames", 1) > 1:
                    raise ValueError("不支持动画图片，请上传静态图片。")
                picture.verify()
        return {
            "mime_type": {"PNG": "image/png", "JPEG": "image/jpeg", "WEBP": "image/webp"}[fmt],
            "data": base64.b64encode(raw).decode("ascii"),
            "filename": Path(filename).name,
        }
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("图片损坏、尺寸过大或格式不受支持。") from exc


def parse_research_file(filename: str, raw: bytes, allow_images: bool = True) -> dict:
    """Parse one upload. Warnings explicitly describe any structural loss."""
    if not filename or not Path(filename).suffix:
        raise ValueError("缺少文件名或文件扩展名。")
    if not raw:
        raise ValueError("文件为空。")
    if len(raw) > MAX_BYTES:
        raise ValueError("文件不能超过 20MB。")
    suffix = Path(filename).suffix.lower()
    if not allow_images and suffix not in REFERENCE_SUFFIXES:
        raise ValueError("参考资料只支持 TXT、Markdown、DOCX 和可提取文字的 PDF。")
    result = {"text": "", "images": [], "kind": suffix.lstrip("."),
              "table_count": 0, "image_count": 0, "warnings": []}

    def add_picture(data: bytes, name: str) -> None:
        result["image_count"] += 1
        if allow_images:
            try:
                result["images"].append(_image(data, name))
            except ValueError as exc:
                result["warnings"].append(f"图片 {Path(name).name} 未提取：{exc}")

    try:
        if suffix in {".txt", ".md", ".csv"}:
            try:
                decoded = raw.decode("utf-8-sig")
            except UnicodeDecodeError as exc:
                raise ValueError("文本文件必须使用 UTF-8 编码。") from exc
            if "\x00" in decoded:
                raise ValueError("文本包含无效的空字符，请确认文件格式和 UTF-8 编码。")
            if suffix == ".csv":
                rows = list(csv.reader(StringIO(decoded)))
                result["text"] = "\n".join("\t".join(row) for row in rows)
                result["table_count"] = 1 if rows else 0
            else:
                result["text"] = decoded

        elif suffix == ".docx":
            _checked_archive(raw)
            document = Document(BytesIO(raw))
            blocks = []
            for element in document.element.body.iterchildren():
                if element.tag.endswith("}p"):
                    paragraph = Paragraph(element, document)
                    content = _paragraph_text(paragraph)
                    if content.strip():
                        blocks.append(content)
                elif element.tag.endswith("}tbl"):
                    table = Table(element, document)
                    result["table_count"] += 1
                    blocks.append("\n".join("\t".join("\n".join(_paragraph_text(p) for p in cell.paragraphs)
                                                     for cell in row.cells) for row in table.rows))
                    if any(cell.tables for row in table.rows for cell in row.cells):
                        result["warnings"].append("DOCX 表格中的嵌套表格未展开，请将关键实验表另传 CSV 或 XLSX。")
            result["text"] = "\n\n".join(blocks)
            for part in document.part.package.parts:
                if str(part.partname).startswith("/word/media/"):
                    add_picture(part.blob, str(part.partname))
                elif str(part.partname).startswith("/word/charts/") and str(part.partname).endswith(".xml"):
                    result["warnings"].append("DOCX 内嵌图表不能直接渲染；实验中的图表请另存为图片上传。")
            if result["image_count"]:
                result["warnings"].append(
                    "DOCX 内嵌图片已单独提取，图片与正文的精确排版关系不保留。" if allow_images
                    else "参考资料仅提取文字与表格，内嵌图片不用于写作分析。")
            if any(header.paragraphs and any(p.text.strip() for p in header.paragraphs)
                   for section in document.sections for header in (section.header, section.footer)):
                result["warnings"].append("DOCX 页眉和页脚未作为正文提取。")
            xml = document.element.xml
            if "oMath" in xml:
                result["warnings"].append("DOCX 公式只提取符号文字，分式、上下标和矩阵布局可能不完整；参考写法仅依据已提取文字，不能作为精确公式依据。")
            if "OLEObject" in xml or "<w:object" in xml:
                result["warnings"].append("DOCX 含旧式嵌入对象（可能是 MathType 公式）；对象中的公式和符号未转为文字，图片数可能包含公式预览，参考写法仅依据已提取文字。")
            if "txbxContent" in xml or "footnoteReference" in xml or "endnoteReference" in xml:
                result["warnings"].append("DOCX 文本框、脚注或尾注可能未完整提取，请人工检查。")

        elif suffix == ".pdf":
            reader = PdfReader(BytesIO(raw))
            if reader.is_encrypted:
                raise ValueError("PDF 已加密，请先解密后上传。")
            if len(reader.pages) > 100:
                raise ValueError("PDF 不能超过 100 页，请拆分后上传。")
            pages = []
            seen_images = set()
            for index, page in enumerate(reader.pages, 1):
                page_text = page.extract_text(extraction_mode="layout") or ""
                if page_text.strip():
                    pages.append(f"[第 {index} 页]\n{page_text}")
                else:
                    result["warnings"].append(f"PDF 第 {index} 页没有可提取文字，可能是扫描页。")
                try:
                    if allow_images:
                        for page_image in page.images:
                            identity = str(getattr(page_image, "indirect_reference", "")) or (page_image.name, len(page_image.data))
                            if identity in seen_images:
                                continue
                            seen_images.add(identity)
                            add_picture(page_image.data, f"page-{index}-{page_image.name}")
                    else:
                        # Counting names avoids decoding images in reference documents.
                        result["image_count"] += len(page.images.keys())
                except Exception:
                    result["warnings"].append(f"PDF 第 {index} 页的部分图片无法提取，请另传图片。")
            result["text"] = "\n\n".join(pages)
            if not result["text"].strip():
                raise ValueError("PDF 未提取到文字，扫描件暂不支持 OCR；请单独上传实验图片，参考资料需提供文字版。")
            result["warnings"].append("PDF 表格按页面文字提取，列结构、阅读顺序和图中文字可能不完整；关键表格建议另传 CSV 或 XLSX。")
            if not allow_images and result["image_count"]:
                result["warnings"].append("参考 PDF 中的图片只计数，不用于写作分析。")

        elif suffix == ".xlsx":
            _checked_archive(raw)
            try:
                from openpyxl import load_workbook
            except ImportError as exc:
                raise ValueError("Excel 解析需要 openpyxl，请安装项目依赖后重试。") from exc
            workbook = load_workbook(BytesIO(raw), read_only=False, data_only=True)
            blocks = []
            try:
                for sheet in workbook.worksheets:
                    if sheet.max_row * sheet.max_column > 1_000_000:
                        raise ValueError("Excel 工作表范围过大，请删除多余空白行列后重试。")
                    rows = ["\t".join("" if cell is None else str(cell) for cell in row)
                            for row in sheet.iter_rows(values_only=True)]
                    rows = [row for row in rows if row.strip()]
                    if rows:
                        blocks.append(f"[工作表：{sheet.title}]\n" + "\n".join(rows))
                        result["table_count"] += 1
                    for index, picture in enumerate(getattr(sheet, "_images", []), 1):
                        try:
                            add_picture(picture._data(), f"{sheet.title}-{index}.{picture.format}")
                        except Exception:
                            result["warnings"].append(f"工作表 {sheet.title} 的图片无法读取，请单独上传。")
                    if getattr(sheet, "_charts", []):
                        result["warnings"].append(f"工作表 {sheet.title} 的原生图表未渲染，已保留表格数据，请另传图表图片。")
            finally:
                workbook.close()
            result["text"] = "\n\n".join(blocks)
            result["warnings"].append("Excel 仅读取单元格缓存值；未缓存的公式可能为空，合并单元格、格式和绘图布局不保留。")

        elif suffix in IMAGE_SUFFIXES:
            result["images"] = [_image(raw, filename)]
            result["image_count"] = 1
            result["kind"] = "image"
        else:
            raise ValueError("不支持此格式。实验附件支持 TXT、MD、CSV、DOCX、PDF、XLSX、PNG、JPEG 和 WEBP。")
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"无法解析 {suffix.lstrip('.').upper()} 文件，文件可能损坏或格式不符。") from exc

    if len(result["text"]) > MAX_CHARS:
        raise ValueError(f"提取文字超过 {MAX_CHARS} 字符，请拆分文件后上传。")
    if not result["text"].strip() and not result["images"]:
        raise ValueError("文件没有可读取的文字或可用图片。")
    result["warnings"] = list(dict.fromkeys(result["warnings"]))
    return result
