"""独立于术语库的写作资料管理与异步文档生成接口。"""
from io import BytesIO
from pathlib import Path
import re
from urllib.parse import quote

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from fastapi import APIRouter, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import StreamingResponse
from starlette.concurrency import run_in_threadpool

from .generation import GenerationInputError, GenerationRequest, validate_generation_inputs
from .research_documents import MAX_BYTES


router = APIRouter()


def public_job(job: dict) -> dict:
    return {key: value for key, value in job.items() if key not in ("request", "references", "assets")}


async def read_upload(file: UploadFile) -> bytes:
    raw = await file.read(MAX_BYTES + 1)
    if len(raw) > MAX_BYTES:
        raise HTTPException(status_code=413, detail="文件不能超过 20MB，请压缩或拆分后上传。")
    return raw


@router.get("/api/references")
def list_references(request: Request):
    return {"references": request.app.state.research.list_references()}


@router.post("/api/references", status_code=201)
async def add_reference(request: Request, title: str = Form("", max_length=180), file: UploadFile = File(...)):
    raw = await read_upload(file)
    filename = file.filename or ""
    try:
        return await run_in_threadpool(request.app.state.research.add_reference,
                                       title.strip() or Path(filename).stem, filename, raw)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@router.get("/api/references/{reference_id}")
def get_reference(request: Request, reference_id: str):
    try:
        return request.app.state.research.get_reference(reference_id)
    except KeyError:
        raise HTTPException(status_code=404, detail="参考资料不存在或已删除。")


@router.delete("/api/references/{reference_id}")
def delete_reference(request: Request, reference_id: str):
    try:
        request.app.state.research.delete_reference(reference_id)
    except KeyError:
        raise HTTPException(status_code=404, detail="参考资料不存在或已删除。")
    return {"deleted": reference_id}


@router.post("/api/generation/assets", status_code=201)
async def add_asset(request: Request, file: UploadFile = File(...)):
    raw = await read_upload(file)
    try:
        return await run_in_threadpool(request.app.state.research.add_asset, file.filename or "", raw)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@router.get("/api/generations")
def list_generations(request: Request):
    return {"generations": [public_job(job) for job in request.app.state.research.list_generations()]}


@router.post("/api/generations", status_code=202)
def start_generation(request: Request, options: GenerationRequest):
    if request.app.state.effective_mode != "llm":
        raise HTTPException(status_code=409, detail="文档生成需要真实模型，请设置 AGENT_MODE=llm 并配置模型 API 后重启服务。")
    store = request.app.state.research
    try:
        references = [store.get_reference(identity) for identity in options.reference_ids]
        assets = [store.get_asset(identity) for identity in options.asset_ids]
    except KeyError:
        raise HTTPException(status_code=409, detail="所选资料或附件已不存在，请刷新后重新选择。")
    try:
        validate_generation_inputs(options, references, assets)
    except GenerationInputError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    job = store.create_generation(options.model_dump(), references, assets)
    request.app.state.generator.start(job["id"])
    return public_job(job)


@router.get("/api/generations/{generation_id}")
def get_generation(request: Request, generation_id: str):
    try:
        return public_job(request.app.state.research.get_generation(generation_id))
    except KeyError:
        raise HTTPException(status_code=404, detail="生成任务不存在。")


@router.get("/api/generations/{generation_id}/download.docx")
def download_generation(request: Request, generation_id: str):
    try:
        result = request.app.state.research.get_generation(generation_id)
    except KeyError:
        raise HTTPException(status_code=404, detail="生成任务不存在。")
    if result["status"] != "completed":
        raise HTTPException(status_code=409, detail="生成尚未完成，无法下载。")
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    document.add_heading(result["title"], 0)
    for line in result["body"].splitlines():
        line = line.strip()
        if not line:
            continue
        if len(line) < 100 and re.match(r"^(#{1,3}\s+|[一二三四五六七八九十]+[、．.]|\d+(?:\.\d+)*[、．. ]\s*)", line):
            document.add_heading(re.sub(r"^#+\s*", "", line), level=1)
        else:
            document.add_paragraph(line)
    document.add_heading("生成与来源说明", level=1)
    document.add_paragraph(f'目标字数：{result["target_words"]}；正文实际：{result["word_count"]}（按非空白字符含标点统计）。')
    document.add_paragraph("下列参考资料仅用于借鉴写作结构与表达；实验事实来源于本次输入与附件。生成内容仍需作者核验。")
    for reference in result.get("reference_sources", []):
        document.add_paragraph(f'写作参考：{reference["title"]}（{reference["filename"]}）')
    if result["request"].get("experiment_text", "").strip():
        document.add_paragraph("实验来源：本次手动输入的实验结果")
    for asset in result.get("asset_sources", []):
        document.add_paragraph(f'实验附件：{asset["filename"]}')
    if result.get("warnings"):
        document.add_heading("待核实与限制", level=1)
        for warning in result["warnings"]:
            document.add_paragraph(warning)
    stream = BytesIO()
    document.save(stream)
    stream.seek(0)
    filename = quote(result["title"][:80] + ".docx", safe="")
    return StreamingResponse(stream, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename=analysis.docx; filename*=UTF-8''{filename}"})
