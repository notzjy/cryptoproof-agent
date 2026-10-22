"""CryptoProof Agent API. 本地演示应用；公网部署前必须补充登录和访问控制。"""
from __future__ import annotations
from contextlib import asynccontextmanager
from io import BytesIO
from pathlib import Path
import os
import sqlite3
import uuid

from fastapi import FastAPI, HTTPException, UploadFile, File, Form
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from starlette.concurrency import run_in_threadpool
from pydantic import BaseModel, Field
from docx import Document
from langgraph.checkpoint.sqlite import SqliteSaver

from .agent import ProofreadingAgent
from .documents import extract_text, MAX_CHARS, MAX_BYTES
from .knowledge import search_term_candidates
from .libraries import LibraryStore, MAX_LIBRARY_BYTES
from .research_store import ResearchStore
from .generation import GenerationService
from .generation_api import router as generation_router

ROOT = Path(__file__).resolve().parent.parent


class SubmitRequest(BaseModel):
    text: str = Field(min_length=1, max_length=MAX_CHARS)
    library_id: str | None = Field(default=None, max_length=80)


class ApprovalRequest(BaseModel):
    approved_ids: list[str] = Field(default_factory=list, max_length=200)


class SearchRequest(BaseModel):
    query: str = Field(min_length=1, max_length=600)
    library_id: str | None = Field(default=None, max_length=80)


def create_app(data_dir: str | Path | None = None, mode: str | None = None) -> FastAPI:
    data_path = Path(data_dir or os.getenv("DATA_DIR") or ROOT / "data")
    effective_mode = (mode or os.getenv("AGENT_MODE") or "demo").lower()

    @asynccontextmanager
    async def lifespan(app: FastAPI):
        data_path.mkdir(parents=True, exist_ok=True)
        conn = sqlite3.connect(data_path / "checkpoints.sqlite", check_same_thread=False)
        try:
            app.state.connection = conn
            app.state.libraries = LibraryStore(data_path)
            app.state.engine = ProofreadingAgent(SqliteSaver(conn), mode=effective_mode)
            app.state.effective_mode = effective_mode
            app.state.research = ResearchStore(data_path)
            app.state.research.recover_interrupted()
            app.state.generator = GenerationService(app.state.research)
            yield
        finally:
            if hasattr(app.state, "generator"):
                await run_in_threadpool(app.state.generator.shutdown)
            conn.close()

    app = FastAPI(title="CryptoProof Agent", version="0.1.0", lifespan=lifespan)
    app.mount("/static", StaticFiles(directory=ROOT / "static"), name="static")
    app.include_router(generation_router)

    def engine() -> ProofreadingAgent:
        return app.state.engine

    def selected_library(library_id: str | None):
        try:
            return app.state.libraries.snapshot(library_id)
        except KeyError:
            raise HTTPException(status_code=409, detail="术语库不存在或已删除，请重新选择或添加术语库")

    def result_for(task_id: str) -> dict:
        snapshot = engine().snapshot(task_id)
        if not snapshot or not snapshot.values or not snapshot.values.get("text"):
            raise HTTPException(status_code=404, detail="任务不存在")
        state = snapshot.values
        waiting = "human_review" in snapshot.next
        # 兼容旧版本：已采纳并生成修订稿的任务曾被警告错误地标成 incomplete。
        legacy_reviewed = (not snapshot.next and state.get("status") == "incomplete"
                           and bool(state.get("approved_ids")) and isinstance(state.get("corrected_text"), str))
        review_completed = state.get("review_completed", False) or legacy_reviewed
        result_status = "completed" if legacy_reviewed else state.get("status", "running")
        return {
            "task_id": task_id,
            "status": "awaiting_review" if waiting else result_status,
            "mode": effective_mode,
            "rag_mode": os.getenv("RAG_MODE", "sparse"),
            "original_text": state["text"],
            "corrected_text": state.get("corrected_text"),
            "edits": state.get("edits", []),
            "approved_ids": state.get("approved_ids", []),
            "trace": engine().tool_trace(task_id),
            "summary": state.get("summary", ""),
            "warnings": state.get("warnings", []),
            "analysis_complete": state.get("analysis_complete", True),
            "review_completed": review_completed,
            "library": {k: v for k, v in state.get("library", {}).items() if k != "terms"},
        }

    @app.get("/api/health")
    def health():
        libraries = app.state.libraries.list()
        return {"status": "ok", "mode": effective_mode,
                "rag_mode": os.getenv("RAG_MODE", "sparse"), "agent": "langgraph",
                "term_count": libraries[0]["term_count"] if libraries else 0,
                "library_count": len(libraries), "term_retriever": "char_tfidf_sequence",
                "model": os.getenv("OPENAI_MODEL", "gpt-4o-mini") if effective_mode == "llm" else None}

    @app.get("/api/libraries")
    def list_libraries():
        libraries = app.state.libraries.list()
        return {"libraries": libraries, "default_library_id": libraries[0]["id"] if libraries else None}

    @app.post("/api/libraries", status_code=201)
    async def add_library(name: str = Form(..., max_length=60), domain: str = Form("", max_length=60),
                          file: UploadFile = File(...)):
        raw = await file.read(MAX_LIBRARY_BYTES + 1)
        if len(raw) > MAX_LIBRARY_BYTES:
            raise HTTPException(status_code=413, detail="术语文件不能超过 1MB")
        try:
            return app.state.libraries.add(name, domain, file.filename or "", raw)
        except ValueError as exc:
            raise HTTPException(status_code=422, detail=str(exc)) from exc

    @app.delete("/api/libraries/{library_id}")
    def delete_library(library_id: str):
        try:
            app.state.libraries.delete(library_id)
        except KeyError:
            raise HTTPException(status_code=404, detail="术语库不存在或已删除")
        return {"deleted": library_id}

    @app.post("/api/tasks")
    def submit(request: SubmitRequest):
        if not request.text.strip():
            raise HTTPException(status_code=422, detail="内容不能为空")
        task_id = str(uuid.uuid4())
        library = selected_library(request.library_id)
        try:
            engine().submit(request.text, task_id, library=library)
        except Exception as exc:
            # 接口错误不返回服务端密钥、请求包、文档正文；详细错误只在本地日志中排查。
            raise HTTPException(status_code=502, detail=f"校对流程执行失败：{type(exc).__name__}") from exc
        return result_for(task_id)

    @app.post("/api/tasks/upload")
    async def submit_file(file: UploadFile = File(...), library_id: str | None = Form(None)):
        if file.size and file.size > MAX_BYTES:
            raise HTTPException(status_code=413, detail="文件超出 3MB 限制")
        raw = await file.read(MAX_BYTES + 1)
        try:
            text = extract_text(file.filename or "", raw)
        except ValueError as exc:
            raise HTTPException(status_code=422, detail=str(exc)) from exc
        return await run_in_threadpool(submit, SubmitRequest(text=text, library_id=library_id))

    @app.get("/api/tasks/{task_id}")
    def get_task(task_id: str):
        return result_for(task_id)

    @app.post("/api/tasks/{task_id}/approve")
    def approve(task_id: str, request: ApprovalRequest):
        snapshot = engine().snapshot(task_id)
        if not snapshot or "human_review" not in snapshot.next:
            raise HTTPException(status_code=409, detail="任务不在等待人工审核状态")
        known = {e["id"] for e in snapshot.values.get("edits", [])}
        if not set(request.approved_ids).issubset(known):
            raise HTTPException(status_code=422, detail="包含不属于该任务的修改项")
        try:
            engine().resume(task_id, request.approved_ids)
        except Exception as exc:
            raise HTTPException(status_code=502, detail=f"审核恢复失败：{type(exc).__name__}") from exc
        return result_for(task_id)

    @app.post("/api/knowledge/search")
    def search(request: SearchRequest):
        library = selected_library(request.library_id)
        return {"query": request.query, "library_id": library["id"],
                "results": search_term_candidates(request.query, records=library["terms"])}

    @app.get("/api/tasks/{task_id}/download.docx")
    def download_docx(task_id: str):
        result = result_for(task_id)
        if result["status"] != "completed":
            raise HTTPException(status_code=409, detail="请先完成人工审核")
        doc = Document()
        doc.add_heading("CryptoProof Agent 校对结果", 0)
        doc.add_paragraph(result["corrected_text"] or "")
        if result["warnings"] or not result["analysis_complete"]:
            doc.add_heading("校对范围与保留说明", level=1)
            if not result["analysis_complete"]:
                doc.add_paragraph("本次检查未完整完成，此修订稿仅应用已通过校验并经用户确认的修改，建议继续复核。")
            for warning in result["warnings"]:
                doc.add_paragraph(warning)
        doc.add_heading("审核记录", level=1)
        for edit in result["edits"]:
            label = "通过" if edit["id"] in result["approved_ids"] else "未采纳"
            doc.add_paragraph(
                f'[{label}] {edit["original"]} → {edit["replacement"]}；'
                f'原因：{edit["reason"]}；来源：{edit["source"]}')
        content = BytesIO()
        doc.save(content)
        content.seek(0)
        return StreamingResponse(content,
                                 media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                                 headers={"Content-Disposition":f'attachment; filename="cryptoproof-{task_id[:8]}.docx"'})

    @app.get("/")
    def home():
        return FileResponse(ROOT / "static" / "index.html")

    @app.get("/generate")
    def generate_page():
        return FileResponse(ROOT / "static" / "generate.html")

    return app


app = create_app()
