from io import BytesIO

import pytest
from docx import Document
from docx.oxml import OxmlElement, parse_xml
from openpyxl import Workbook
from PIL import Image
from pypdf import PdfWriter

from cryptoproof.research_documents import MAX_BYTES, MAX_CHARS, parse_research_file
from cryptoproof.research_store import ResearchStore


def docx_bytes():
    document = Document()
    document.add_paragraph("实验开始：采样 100 次。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "算法"
    table.cell(0, 1).text = "成功率"
    table.cell(1, 0).text = "SM4"
    table.cell(1, 1).text = "91%"
    document.add_paragraph("实验结束：结果仅用于本次测试。")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_preserves_paragraph_and_table_order():
    parsed = parse_research_file("results.docx", docx_bytes())
    assert parsed["table_count"] == 1
    assert parsed["text"].index("实验开始") < parsed["text"].index("成功率") < parsed["text"].index("实验结束")
    assert "SM4\t91%" in parsed["text"]


def test_docx_equation_tokens_keep_order_and_warn_about_layout_loss():
    document = Document()
    paragraph = document.add_paragraph("定义：")
    equation = OxmlElement("m:oMath")
    math_run = OxmlElement("m:r")
    token = OxmlElement("m:t")
    token.text = "x+y"
    math_run.append(token)
    equation.append(math_run)
    paragraph._p.append(equation)
    paragraph.add_run("，其中变量如下。")
    buffer = BytesIO()
    document.save(buffer)
    parsed = parse_research_file("formula.docx", buffer.getvalue(), allow_images=False)
    assert "定义：x+y，其中变量如下。" in parsed["text"]
    assert any("公式" in warning and "已提取文字" in warning for warning in parsed["warnings"])


def test_docx_legacy_mathtype_objects_have_explicit_warning():
    document = Document()
    paragraph = document.add_paragraph("旧式公式如下。")
    paragraph._p.append(parse_xml(
        '<w:object xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:o="urn:schemas-microsoft-com:office:office"><o:OLEObject Type="Embed"/></w:object>'
    ))
    buffer = BytesIO()
    document.save(buffer)
    parsed = parse_research_file("legacy.docx", buffer.getvalue(), allow_images=False)
    assert any("MathType" in warning and "未转为文字" in warning for warning in parsed["warnings"])


def test_csv_and_xlsx_keep_experiment_table_values():
    csv = parse_research_file("results.csv", "算法,成功率\nSM4,0.91\n".encode())
    assert csv["table_count"] == 1 and "SM4\t0.91" in csv["text"]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "攻击结果"
    sheet.append(["算法", "成功率"])
    sheet.append(["SM4", 0.91])
    buffer = BytesIO()
    workbook.save(buffer)
    parsed = parse_research_file("results.xlsx", buffer.getvalue())
    assert parsed["table_count"] == 1 and "SM4\t0.91" in parsed["text"]
    assert "攻击结果" in parsed["text"]
    assert any("公式" in warning for warning in parsed["warnings"])


def test_image_validation_and_embedded_reference_image_handling():
    buffer = BytesIO()
    Image.new("RGB", (20, 10), "white").save(buffer, format="PNG")
    image_raw = buffer.getvalue()
    parsed = parse_research_file("plot.png", image_raw)
    assert parsed["kind"] == "image" and parsed["image_count"] == 1
    assert parsed["images"][0]["mime_type"] == "image/png"
    assert parsed["images"][0]["data"]
    document = Document()
    document.add_paragraph("本研究图示如下。")
    document.add_picture(BytesIO(image_raw))
    doc_buffer = BytesIO()
    document.save(doc_buffer)
    reference = parse_research_file("paper.docx", doc_buffer.getvalue(), allow_images=False)
    assert reference["image_count"] == 1 and reference["images"] == []
    assert reference["warnings"]
    experiment = parse_research_file("evidence.docx", doc_buffer.getvalue())
    assert len(experiment["images"]) == 1


def test_reference_crud_is_persistent_and_job_uses_frozen_snapshot(tmp_path):
    store = ResearchStore(tmp_path)
    assert store.list_references() == []
    metadata = store.add_reference("分析论文", "paper.docx", docx_bytes())
    assert metadata["table_count"] == 1 and metadata["char_count"] > 0
    assert "text" not in metadata and "images" not in metadata
    source = store.get_reference(metadata["id"])
    asset_meta = store.add_asset("experiment.txt", "攻击成功率 91%，共 100 次。".encode())
    assert "text" not in asset_meta and "images" not in asset_meta
    evidence = store.get_asset(asset_meta["id"])
    request = {"title": "SM4 攻击分析", "target_words": 800, "experiment_text": "结果如附件"}
    job = store.create_generation(request, [source], [evidence])
    request["title"] = "不应修改快照"
    source["text"] = "不应修改快照"
    store.delete_reference(metadata["id"])
    restarted = ResearchStore(tmp_path)
    assert restarted.list_references() == []
    snapshot = restarted.get_generation(job["id"])
    assert snapshot["request"]["title"] == "SM4 攻击分析"
    assert "成功率" in snapshot["references"][0]["text"]
    assert "91%" in snapshot["assets"][0]["text"]
    assert "text" not in snapshot["reference_sources"][0]
    assert "text" not in snapshot["asset_sources"][0]
    with pytest.raises(KeyError):
        restarted.get_reference(metadata["id"])
    with pytest.raises(KeyError):
        restarted.delete_reference(metadata["id"])
    snapshot = restarted.update_generation(job["id"], status="completed", body="分析正文", word_count=4, within_target=True)
    assert snapshot["body"] == "分析正文"
    assert "references" not in restarted.list_generations()[0]


def test_restart_marks_only_unfinished_jobs_failed(tmp_path):
    store = ResearchStore(tmp_path)
    for status in ("queued", "analyzing", "writing", "revising", "completed", "failed"):
        job = store.create_generation({"title": status}, [], [])
        store.update_generation(job["id"], status=status)
    assert store.recover_interrupted() == 4
    statuses = {job["title"]: job["status"] for job in store.list_generations()}
    assert statuses == {"queued": "failed", "analyzing": "failed", "writing": "failed", "revising": "failed", "completed": "completed", "failed": "failed"}
    assert store.recover_interrupted() == 0


def test_job_list_is_limited_and_unknown_ids_are_errors(tmp_path):
    store = ResearchStore(tmp_path)
    for index in range(22):
        store.create_generation({"title": str(index)}, [], [])
    assert len(store.list_generations()) == 20
    for action in (store.get_generation, store.get_asset):
        with pytest.raises(KeyError):
            action("missing")
    with pytest.raises(KeyError):
        store.update_generation("missing", status="failed")


@pytest.mark.parametrize(("filename", "raw"), [
    ("broken.docx", b"not an office document"), ("broken.xlsx", b"bad zip"),
    ("broken.pdf", b"bad pdf"), ("bad.png", b"not a picture"),
    ("unknown.exe", b"abc"), ("invalid.txt", b"\xff\xfe"),
    ("empty.txt", b"\n "), ("binary.txt", b"x\x00y"),
])
def test_bad_or_unsupported_files_are_explicit_errors(filename, raw):
    with pytest.raises(ValueError):
        parse_research_file(filename, raw)


def test_text_size_and_reference_type_limits(tmp_path):
    with pytest.raises(ValueError, match="20MB"):
        parse_research_file("large.txt", b"a" * (MAX_BYTES + 1))
    with pytest.raises(ValueError, match="180000"):
        parse_research_file("long.txt", b"a" * (MAX_CHARS + 1))
    with pytest.raises(ValueError, match="参考资料"):
        ResearchStore(tmp_path).add_reference("图", "plot.png", b"fake")
    with pytest.raises(ValueError, match="资料名称"):
        ResearchStore(tmp_path).add_reference(" ", "paper.txt", b"abc")


def test_scanned_pdf_requires_text_version_or_separate_images():
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    buffer = BytesIO()
    writer.write(buffer)
    with pytest.raises(ValueError, match="扫描件"):
        parse_research_file("scan.pdf", buffer.getvalue())
