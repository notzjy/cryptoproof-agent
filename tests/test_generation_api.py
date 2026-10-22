"""独立资料库与真实 SQLite 后台写作任务的 API 集成测试；全部模型使用替身。"""
import base64
from io import BytesIO
import json
import time

from docx import Document
from fastapi.testclient import TestClient
from langchain_core.messages import AIMessage
from openpyxl import Workbook
from PIL import Image
import pytest

from cryptoproof.api import create_app
import cryptoproof.generation as generation
from cryptoproof.research_documents import MAX_BYTES


EXPERIMENT = "第17轮实验，rk0四个字节命中1个，查询次数20480。"
LIMITATION = "当前只有一个目标密钥，不能推断多密钥恢复率。"
REFERENCE_TEXT = "本论文按实验条件、观察结果、分析和局限组织内容。参考资料有1000把密钥。"
BODY = "实验结果。" * 60


class UnusedProofreadingModel:
    def bind_tools(self, tools):
        return self

    def invoke(self, messages):
        raise AssertionError("写作 API 测试不应调用校对模型")


class GenerationModel:
    def __init__(self):
        self.calls = []
        self.vision_builds = []
        self.fail = False

    def build(self, vision=False):
        self.vision_builds.append(vision)
        return self

    def invoke(self, messages):
        self.calls.append(messages)
        if self.fail:
            return AIMessage(content="故意返回无效 JSON")
        system = messages[0].content
        if system == generation.FACT_PROMPT:
            blocks = messages[-1].content
            payload = json.loads(blocks[0]["text"] if isinstance(blocks, list) else blocks)
            text = payload["experiment_text"]
            result = {"summary": "第17轮单目标的四个字节实验。",
                      "facts": [{"statement": "四个轮密钥字节命中1个。",
                                 "evidence_quote": text or "1/4"}],
                      "limitations": [LIMITATION], "conflicts": []}
        elif system.startswith(generation.WRITE_PROMPT):
            result = {"title": "实验分析", "body": BODY, "missing_information": []}
        elif system == generation.AUDIT_PROMPT:
            result = {"supported": True, "issues": []}
        else:
            raise AssertionError("未预期的模型提示词")
        return AIMessage(content=json.dumps(result, ensure_ascii=False))


@pytest.fixture
def model(monkeypatch):
    instance = GenerationModel()
    monkeypatch.setenv("OPENAI_API_KEY", "test-only-no-network")
    monkeypatch.setattr("cryptoproof.agent.ChatOpenAI", lambda **kwargs: UnusedProofreadingModel())
    monkeypatch.setattr(generation, "build_generation_model", instance.build)
    return instance


def add_reference(client, title="实验论文", filename="paper.txt", text=REFERENCE_TEXT):
    response = client.post("/api/references", data={"title": title},
                           files={"file": (filename, text.encode("utf-8"), "text/plain")})
    assert response.status_code == 201, response.text
    return response.json()


def start_job(client, reference_id, **overrides):
    return client.post("/api/generations", json={
        "title": "第17轮实验分析", "experiment_text": EXPERIMENT,
        "target_words": 300, "reference_ids": [reference_id], **overrides,
    })


def await_job(client, generation_id):
    deadline = time.monotonic() + 2
    while time.monotonic() < deadline:
        response = client.get(f"/api/generations/{generation_id}")
        assert response.status_code == 200, response.text
        job = response.json()
        if job["status"] in {"completed", "failed"}:
            return job
        time.sleep(.01)
    pytest.fail("模型替身后台任务未在2秒内完成")


def png_bytes():
    stream = BytesIO()
    Image.new("RGB", (8, 6), "white").save(stream, format="PNG")
    return stream.getvalue()


def xlsx_bytes():
    book = Workbook()
    book.active.append(["rounds", "top1"])
    book.active.append([17, "1/4"])
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def assert_public_job(job, image_data=None):
    assert not {"request", "references", "assets"} & job.keys()
    for source in [*job.get("reference_sources", []), *job.get("asset_sources", [])]:
        assert "text" not in source and "images" not in source
    if image_data:
        assert image_data not in json.dumps(job, ensure_ascii=False)


def test_reference_crud_is_persistent_and_separate_from_terminology(tmp_path, model):
    data_dir = tmp_path / "data"
    with TestClient(create_app(data_dir, mode="demo")) as client:
        before = client.get("/api/libraries").json()
        assert client.get("/api/references").json() == {"references": []}
        reference = add_reference(client)
        identity = reference["id"]
        assert reference["title"] == "实验论文" and reference["filename"] == "paper.txt"
        assert reference["char_count"] == len(REFERENCE_TEXT)
        assert "text" not in reference
        selected = client.get(f"/api/references/{identity}")
        assert selected.status_code == 200
        assert selected.json()["text"] == REFERENCE_TEXT
        assert client.get("/api/libraries").json() == before
        assert client.get("/api/references").json()["references"] == [reference]
    with TestClient(create_app(data_dir, mode="demo")) as client:
        assert client.get(f"/api/references/{identity}").json()["text"] == REFERENCE_TEXT
        assert client.delete(f"/api/references/{identity}").json() == {"deleted": identity}
        assert client.get("/api/references").json() == {"references": []}
        assert client.get("/api/libraries").json() == before
        assert client.get(f"/api/references/{identity}").status_code == 404
        assert client.delete(f"/api/references/{identity}").status_code == 404


@pytest.mark.parametrize("route", ["/api/references", "/api/generation/assets"])
def test_twenty_megabyte_upload_limit_is_enforced(tmp_path, model, route):
    with TestClient(create_app(tmp_path / "data", mode="demo")) as client:
        response = client.post(route, data={"title": "超大资料"},
                               files={"file": ("oversize.txt", b"x" * (MAX_BYTES + 1), "text/plain")})
        assert response.status_code == 413
        assert "20MB" in response.json()["detail"]
        assert client.get("/api/references").json()["references"] == []


@pytest.mark.parametrize("route,filename,content", [
    ("/api/references", "bad.exe", b"not a reference"),
    ("/api/references", "broken.docx", b"not a docx"),
    ("/api/references", "empty.txt", b""),
    ("/api/references", "bad.txt", b"\xff\xfe\x00"),
    ("/api/generation/assets", "broken.png", b"not a png"),
    ("/api/generation/assets", "bad.exe", b"not an experiment"),
])
def test_bad_uploads_return_validation_errors(tmp_path, model, route, filename, content):
    with TestClient(create_app(tmp_path / "data", mode="demo")) as client:
        response = client.post(route, data={"title": "错误资料"},
                               files={"file": (filename, content, "application/octet-stream")})
        assert response.status_code == 422, response.text
        assert response.json()["detail"]


def test_experiment_image_csv_and_xlsx_upload_preserve_evidence_without_public_base64(tmp_path, model):
    with TestClient(create_app(tmp_path / "data", mode="llm")) as client:
        image = png_bytes()
        uploaded = client.post("/api/generation/assets", files={"file": ("plot.png", image, "image/png")})
        assert uploaded.status_code == 201, uploaded.text
        asset = uploaded.json()
        assert asset["kind"] == "image" and asset["image_count"] == 1
        assert "images" not in asset and "text" not in asset
        internal = client.app.state.research.get_asset(asset["id"])
        assert internal["images"][0]["data"] == base64.b64encode(image).decode("ascii")

        for filename, raw, kind in [
            ("result.csv", b"rounds,top1\n17,1/4\n", "csv"),
            ("result.xlsx", xlsx_bytes(), "xlsx"),
        ]:
            response = client.post("/api/generation/assets", files={"file": (filename, raw)})
            assert response.status_code == 201, response.text
            parsed = response.json()
            assert parsed["table_count"] == 1 and parsed["kind"] == kind
            text = client.app.state.research.get_asset(parsed["id"])["text"]
            assert "17\t1/4" in text


def test_demo_mode_rejects_generation_and_does_not_call_model(tmp_path, model):
    with TestClient(create_app(tmp_path / "data", mode="demo")) as client:
        reference = add_reference(client)
        response = start_job(client, reference["id"])
        assert response.status_code == 409
        assert "AGENT_MODE=llm" in response.json()["detail"]
        assert model.calls == []
        assert client.get("/api/generations").json()["generations"] == []


def test_llm_generation_completes_and_exports_real_docx_with_sources_and_limits(tmp_path, model):
    with TestClient(create_app(tmp_path / "data", mode="llm")) as client:
        reference = add_reference(client)
        response = start_job(client, reference["id"])
        assert response.status_code == 202, response.text
        submitted = response.json()
        assert submitted["status"] == "queued"
        assert_public_job(submitted)
        job = await_job(client, submitted["id"])
        assert job["status"] == "completed", job
        assert job["body"] == BODY and job["word_count"] == 300 and job["within_target"]
        assert job["reference_sources"][0]["id"] == reference["id"]
        assert LIMITATION in job["warnings"]
        assert_public_job(job)
        assert len(model.calls) == 3
        downloaded = client.get(f'/api/generations/{job["id"]}/download.docx')
        assert downloaded.status_code == 200
        assert downloaded.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        assert "attachment;" in downloaded.headers["content-disposition"]
        assert downloaded.content.startswith(b"PK")
        document = Document(BytesIO(downloaded.content))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        for required in (job["title"], BODY, "生成与来源说明", "实验论文", "paper.txt",
                         "本次手动输入的实验结果", "正文实际：300", "待核实与限制", LIMITATION):
            assert required in text


@pytest.mark.parametrize("field", ["reference_ids", "asset_ids"])
def test_unknown_selected_source_returns_conflict(tmp_path, model, field):
    with TestClient(create_app(tmp_path / "data", mode="llm")) as client:
        reference = add_reference(client)
        response = start_job(client, reference["id"], **{field: ["missing-source"]})
        assert response.status_code == 409
        assert "已不存在" in response.json()["detail"]
        assert model.calls == []


@pytest.mark.parametrize("field", ["reference_ids", "asset_ids"])
def test_duplicate_selected_sources_are_rejected(tmp_path, model, field):
    with TestClient(create_app(tmp_path / "data", mode="llm")) as client:
        reference = add_reference(client)
        if field == "reference_ids":
            identity = reference["id"]
        else:
            response = client.post("/api/generation/assets", files={"file": ("result.txt", EXPERIMENT.encode("utf-8"))})
            assert response.status_code == 201
            identity = response.json()["id"]
        response = start_job(client, reference["id"], **{field: [identity, identity]})
        assert response.status_code == 422
        assert "重复" in response.json()["detail"]
        assert model.calls == []


def test_failed_job_is_publicly_reported_and_cannot_be_downloaded(tmp_path, model):
    model.fail = True
    with TestClient(create_app(tmp_path / "data", mode="llm")) as client:
        reference = add_reference(client)
        response = start_job(client, reference["id"])
        assert response.status_code == 202
        identity = response.json()["id"]
        job = await_job(client, identity)
        assert job["status"] == "failed" and job["error"]
        assert job["body"] == ""
        assert_public_job(job)
        response = client.get(f"/api/generations/{identity}/download.docx")
        assert response.status_code == 409
        assert client.get("/api/generations/unknown").status_code == 404
        assert client.get("/api/generations/unknown/download.docx").status_code == 404


def test_reference_deletion_keeps_historical_snapshots_results_and_downloads(tmp_path, model):
    data_dir = tmp_path / "data"
    with TestClient(create_app(data_dir, mode="llm")) as client:
        reference = add_reference(client)
        identity = start_job(client, reference["id"]).json()["id"]
        completed = await_job(client, identity)
        assert completed["status"] == "completed"
        assert client.delete(f'/api/references/{reference["id"]}').status_code == 200
        assert client.get(f'/api/references/{reference["id"]}').status_code == 404
        assert client.get("/api/references").json()["references"] == []
        internal = client.app.state.research.get_generation(identity)
        assert internal["references"][0]["text"] == REFERENCE_TEXT
        assert internal["request"]["reference_ids"] == [reference["id"]]
        history = client.get("/api/generations").json()["generations"]
        assert history == [completed]
        assert_public_job(history[0])
    with TestClient(create_app(data_dir, mode="llm")) as client:
        job = client.get(f"/api/generations/{identity}").json()
        assert job == completed
        downloaded = client.get(f"/api/generations/{identity}/download.docx")
        assert downloaded.status_code == 200
        text = "\n".join(p.text for p in Document(BytesIO(downloaded.content)).paragraphs)
        assert "实验论文" in text and "paper.txt" in text and BODY in text
        assert client.get("/api/generations").json()["generations"][0]["id"] == identity


def test_public_jobs_hide_inputs_and_raw_images_in_submit_poll_and_history(tmp_path, model):
    with TestClient(create_app(tmp_path / "data", mode="llm")) as client:
        reference = add_reference(client)
        raw = png_bytes()
        image_data = base64.b64encode(raw).decode("ascii")
        uploaded = client.post("/api/generation/assets", files={"file": ("plot.png", raw, "image/png")})
        assert uploaded.status_code == 201
        response = start_job(client, reference["id"], experiment_text="", asset_ids=[uploaded.json()["id"]],
                             instructions="PRIVATE_REQUEST_SENTINEL")
        assert response.status_code == 202
        assert_public_job(response.json(), image_data)
        identity = response.json()["id"]
        job = await_job(client, identity)
        assert job["status"] == "completed", job
        assert_public_job(job, image_data)
        assert "PRIVATE_REQUEST_SENTINEL" not in json.dumps(job, ensure_ascii=False)
        history = client.get("/api/generations").json()["generations"]
        assert len(history) == 1
        assert_public_job(history[0], image_data)
        assert model.vision_builds == [True, False]
        image_blocks = model.calls[0][-1].content
        assert any(block.get("type") == "image_url" and image_data in block["image_url"]["url"]
                   for block in image_blocks)
